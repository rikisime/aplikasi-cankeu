import io
from datetime import date
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from database import db
from models import User
from auth import get_current_user
from routes.reports import daterange_labels

router = APIRouter(prefix="/export", tags=["export"])


def format_rp(value: float) -> str:
    return f"Rp {value:,.0f}".replace(",", ".")


async def get_period_transactions(user_id: str, period: str, anchor_date: str = None):
    anchor = date.fromisoformat(anchor_date) if anchor_date else date.today()
    start, end, _, _ = daterange_labels(period, anchor)
    txs = await db.transactions.find(
        {"user_id": user_id, "date": {"$gte": start.isoformat(), "$lte": end.isoformat()}},
        {"_id": 0},
    ).sort("date", 1).to_list(10000)
    return txs, start, end


@router.get("/excel")
async def export_excel(period: str = "monthly", anchor_date: str = None, user: User = Depends(get_current_user)):
    txs, start, end = await get_period_transactions(user.user_id, period, anchor_date)

    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Keuangan"

    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    ws.append([f"Laporan Keuangan ArthaKu ({start.isoformat()} - {end.isoformat()})"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Pengguna: {user.name}"])
    ws.append([])

    headers = ["Tanggal", "Tipe", "Kategori", "Deskripsi", "Metode", "Jumlah (Rp)"]
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    total_income = 0.0
    total_expense = 0.0
    for tx in txs:
        ws.append([
            tx["date"],
            "Pemasukan" if tx["type"] == "income" else "Pengeluaran",
            tx["category_name"],
            tx.get("description", ""),
            tx.get("payment_method", ""),
            tx["amount"],
        ])
        if tx["type"] == "income":
            total_income += tx["amount"]
        else:
            total_expense += tx["amount"]

    ws.append([])
    ws.append(["", "", "", "", "Total Pemasukan", total_income])
    ws.append(["", "", "", "", "Total Pengeluaran", total_expense])
    ws.append(["", "", "", "", "Saldo Bersih", total_income - total_expense])

    for col, width in zip("ABCDEF", [12, 12, 22, 30, 14, 16]):
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"laporan-keuangan-{period}-{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/pdf")
async def export_pdf(period: str = "monthly", anchor_date: str = None, user: User = Depends(get_current_user)):
    txs, start, end = await get_period_transactions(user.user_id, period, anchor_date)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title2", parent=styles["Title"], textColor=colors.HexColor("#2563EB"))
    elements = []

    elements.append(Paragraph("ArthaKu - Laporan Keuangan", title_style))
    elements.append(Paragraph(f"Periode: {start.isoformat()} s/d {end.isoformat()}", styles["Normal"]))
    elements.append(Paragraph(f"Pengguna: {user.name} ({user.email})", styles["Normal"]))
    elements.append(Spacer(1, 16))

    total_income = sum(t["amount"] for t in txs if t["type"] == "income")
    total_expense = sum(t["amount"] for t in txs if t["type"] == "expense")
    net = total_income - total_expense

    summary_data = [
        ["Total Pemasukan", format_rp(total_income)],
        ["Total Pengeluaran", format_rp(total_expense)],
        ["Saldo Bersih", format_rp(net)],
    ]
    summary_table = Table(summary_data, colWidths=[8 * cm, 8 * cm])
    summary_table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (1, 0), (1, 0), colors.HexColor("#10B981")),
        ("TEXTCOLOR", (1, 1), (1, 1), colors.HexColor("#EF4444")),
        ("TEXTCOLOR", (1, 2), (1, 2), colors.HexColor("#2563EB")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 20))

    table_data = [["Tanggal", "Tipe", "Kategori", "Deskripsi", "Jumlah"]]
    for tx in txs:
        table_data.append([
            tx["date"],
            "Masuk" if tx["type"] == "income" else "Keluar",
            tx["category_name"],
            (tx.get("description") or "-")[:30],
            format_rp(tx["amount"]),
        ])

    tx_table = Table(table_data, colWidths=[2.3 * cm, 1.8 * cm, 4 * cm, 5.4 * cm, 3.5 * cm], repeatRows=1)
    tx_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elements.append(tx_table)

    doc.build(elements)
    buffer.seek(0)
    filename = f"laporan-keuangan-{period}-{date.today().isoformat()}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
